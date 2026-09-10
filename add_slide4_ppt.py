import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def add_slide4():
    pptx_path = r"d:\PROJECTS-PUPU\satquery-ai\SatQuery_AI_SIH26167_Presentation.pptx"
    prs = Presentation(pptx_path)
    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(11, 44, 91)
    ACCENT_BLUE = RGBColor(24, 119, 242)
    DARK_GRAY = RGBColor(33, 37, 41)
    LIGHT_BG = RGBColor(248, 249, 250)
    CARD_BORDER = RGBColor(220, 224, 230)
    WHITE = RGBColor(255, 255, 255)

    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY & CHALLENGES
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)

    # Top Left Oval (Team Name)
    oval = s4.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(0.4), Inches(2.2), Inches(1.1))
    oval.fill.solid(); oval.fill.fore_color.rgb = WHITE
    oval.line.color.rgb = RGBColor(128, 0, 128); oval.line.width = Pt(2.0)
    otf = oval.text_frame; otf.text = "Your Team\nName"
    otf.paragraphs[0].font.size = Pt(14); otf.paragraphs[0].font.bold = True
    otf.paragraphs[0].font.color.rgb = DARK_GRAY; otf.paragraphs[0].alignment = PP_ALIGN.CENTER
    otf.paragraphs[1].font.size = Pt(14); otf.paragraphs[1].font.bold = True
    otf.paragraphs[1].font.color.rgb = DARK_GRAY; otf.paragraphs[1].alignment = PP_ALIGN.CENTER

    # Top Center Title: FEASIBILITY & CHALLENGES
    itb = s4.shapes.add_textbox(Inches(3.2), Inches(0.5), Inches(6.8), Inches(0.9))
    itf = itb.text_frame; itf.word_wrap = True
    ip = itf.paragraphs[0]
    ip.text = "FEASIBILITY & CHALLENGES"
    ip.font.size = Pt(28); ip.font.bold = True; ip.font.color.rgb = DARK_GRAY
    ip.alignment = PP_ALIGN.CENTER

    # Top Right: SIH 2026 Badge
    sih_box = s4.shapes.add_textbox(Inches(10.2), Inches(0.4), Inches(2.5), Inches(1.1))
    stf = sih_box.text_frame
    sp = stf.paragraphs[0]
    sp.text = "SMART INDIA\nHACKATHON 2026"
    sp.font.size = Pt(15); sp.font.bold = True; sp.font.color.rgb = NAVY
    sp.alignment = PP_ALIGN.RIGHT

    # 3 Cards Grid: Feasibility (Left), Challenges (Center), Strategies (Right)
    cards = [
        ("Analysis of the feasibility of the idea", [
            ("Technical: ", "Proven working end-to-end prototype integrating GeoTIFF ingestion, adapted RS-VLM, and React MapLibre GIS."),
            ("Operational: ", "Self-contained with zero proprietary API dependencies; runs portably on standard CPU hardware."),
            ("Economic: ", "Built on open-source libraries (GDAL/PyTorch), eliminating expensive per-seat GIS software licensing costs.")
        ], Inches(0.6), Inches(3.9)),
        ("Potential challenges and risks", [
            ("Data Heterogeneity: ", "Varying Coordinate Reference Systems (CRS), spatial resolution (GSD), and band orders across sensors."),
            ("Atmospheric & Noise: ", "Optical scenes suffer cloud cover and shadowing; SAR imagery contains granular speckle noise."),
            ("AI Hallucination: ", "Standard vision-language models risk hallucinating ungrounded answers for indistinct satellite features.")
        ], Inches(4.7), Inches(3.9)),
        ("Strategies for overcoming these challenges", [
            ("Automated Ingestion: ", "RasterIngestor auto-aligns CRS, resamples pixels, and calibrates SAR σ⁰ dB and optical reflectance."),
            ("Optical + SAR Synergy: ", "Fuses cloud-penetrating SAR radar backscatter with optical spectral context, filtering radar speckle."),
            ("Evidence & Audit: ", "Pairs all responses with spatial GeoJSON bounding boxes/masks, honest confidence, and immutable audit logs.")
        ], Inches(8.8), Inches(3.9))
    ]

    for title, bullets, left_pos, width in cards:
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.8), width, Inches(4.9))
        card.fill.solid(); card.fill.fore_color.rgb = LIGHT_BG
        card.line.color.rgb = CARD_BORDER; card.line.width = Pt(1.5)

        ch_box = s4.shapes.add_textbox(left_pos + Inches(0.15), Inches(1.9), width - Inches(0.3), Inches(0.6))
        ch_p = ch_box.text_frame.paragraphs[0]
        ch_p.text = "• " + title
        ch_p.font.size = Pt(15); ch_p.font.bold = True; ch_p.font.color.rgb = NAVY

        cb_box = s4.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.5), width - Inches(0.3), Inches(4.0))
        cb_tf = cb_box.text_frame; cb_tf.word_wrap = True

        for idx, (lead, desc) in enumerate(bullets):
            p = cb_tf.paragraphs[0] if idx == 0 else cb_tf.add_paragraph()
            p.space_after = Pt(10)
            r1 = p.add_run(); r1.text = "• " + lead; r1.font.bold = True; r1.font.size = Pt(11.5); r1.font.color.rgb = DARK_GRAY
            r2 = p.add_run(); r2.text = desc; r2.font.bold = False; r2.font.size = Pt(11); r2.font.color.rgb = DARK_GRAY

    # Footer
    fbar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), Inches(13.333), Inches(0.6))
    fbar.fill.solid(); fbar.fill.fore_color.rgb = ACCENT_BLUE; fbar.line.fill.background()

    ftxt_l = s4.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(8.0), Inches(0.5))
    ftxt_l.text_frame.text = "@SIH Idea submission- Template"
    ftxt_l.text_frame.paragraphs[0].font.color.rgb = WHITE
    ftxt_l.text_frame.paragraphs[0].font.size = Pt(12)

    ftxt_r = s4.shapes.add_textbox(Inches(12.0), Inches(6.95), Inches(1.0), Inches(0.5))
    ftxt_r.text_frame.text = "4"
    ftxt_r.text_frame.paragraphs[0].font.color.rgb = WHITE
    ftxt_r.text_frame.paragraphs[0].font.size = Pt(14)
    ftxt_r.text_frame.paragraphs[0].font.bold = True

    prs.save(pptx_path)
    print("Slide 4 appended successfully!")

if __name__ == "__main__":
    add_slide4()
