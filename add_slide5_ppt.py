import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def add_slide5():
    pptx_path = r"d:\PROJECTS-PUPU\satquery-ai\SatQuery_AI_SIH26167_Presentation.pptx"
    prs = Presentation(pptx_path)
    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(11, 44, 91)
    ACCENT_BLUE = RGBColor(24, 119, 242)
    DARK_GRAY = RGBColor(33, 37, 41)
    LIGHT_BG = RGBColor(248, 249, 250)
    CARD_BORDER = RGBColor(220, 224, 230)
    WHITE = RGBColor(255, 255, 255)

    # -------------------------------------------------------------
    # SLIDE 5: IMPACT & BENEFITS
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)

    # Top Left Oval (Team Name)
    oval = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(0.4), Inches(2.2), Inches(1.1))
    oval.fill.solid(); oval.fill.fore_color.rgb = WHITE
    oval.line.color.rgb = RGBColor(128, 0, 128); oval.line.width = Pt(2.0)
    otf = oval.text_frame; otf.text = "Your Team\nName"
    otf.paragraphs[0].font.size = Pt(14); otf.paragraphs[0].font.bold = True
    otf.paragraphs[0].font.color.rgb = DARK_GRAY; otf.paragraphs[0].alignment = PP_ALIGN.CENTER
    otf.paragraphs[1].font.size = Pt(14); otf.paragraphs[1].font.bold = True
    otf.paragraphs[1].font.color.rgb = DARK_GRAY; otf.paragraphs[1].alignment = PP_ALIGN.CENTER

    # Top Center Title: IMPACT & BENEFITS
    itb = s5.shapes.add_textbox(Inches(3.2), Inches(0.5), Inches(6.8), Inches(0.9))
    itf = itb.text_frame; itf.word_wrap = True
    ip = itf.paragraphs[0]
    ip.text = "IMPACT & BENEFITS"
    ip.font.size = Pt(30); ip.font.bold = True; ip.font.color.rgb = DARK_GRAY
    ip.alignment = PP_ALIGN.CENTER

    # Top Right: SIH 2026 Badge
    sih_box = s5.shapes.add_textbox(Inches(10.2), Inches(0.4), Inches(2.5), Inches(1.1))
    stf = sih_box.text_frame
    sp = stf.paragraphs[0]
    sp.text = "SMART INDIA\nHACKATHON 2026"
    sp.font.size = Pt(15); sp.font.bold = True; sp.font.color.rgb = NAVY
    sp.alignment = PP_ALIGN.RIGHT

    # 2 Cards: Left = Target Audience Impact (5.9 in), Right = Multi-Sector Benefits (5.9 in)
    c1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.8), Inches(5.9), Inches(4.9))
    c1.fill.solid(); c1.fill.fore_color.rgb = LIGHT_BG
    c1.line.color.rgb = CARD_BORDER; c1.line.width = Pt(1.5)

    c1_hdr = s5.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(5.5), Inches(0.6))
    c1_hp = c1_hdr.text_frame.paragraphs[0]
    c1_hp.text = "• Potential impact on the target audience"
    c1_hp.font.size = Pt(15); c1_hp.font.bold = True; c1_hp.font.color.rgb = NAVY

    c1_body = s5.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(5.5), Inches(4.1))
    c1_btf = c1_body.text_frame; c1_btf.word_wrap = True

    impact_items = [
        ("Disaster Response (NDRF/SDMA): ", "Instant flood extent boundaries & affected area stats (<70ms) to prioritize rescues."),
        ("Water Authorities (Jal Shakti): ", "Automates reservoir tracking and wetland depletion monitoring without GIS staff."),
        ("Urban Municipalities: ", "Detects unauthorized built-up sprawl and green cover loss via bi-temporal change maps."),
        ("District Administrators: ", "Enables non-technical field officers to query satellite scenes in plain natural language."),
        ("Agriculture & Forestry: ", "Tracks crop canopy stress, seasonal vegetation health, and illegal deforestation trends.")
    ]

    for idx, (lead, desc) in enumerate(impact_items):
        p = c1_btf.paragraphs[0] if idx == 0 else c1_btf.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run(); r1.text = "• " + lead; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = DARK_GRAY
        r2 = p.add_run(); r2.text = desc; r2.font.bold = False; r2.font.size = Pt(10.5); r2.font.color.rgb = DARK_GRAY

    # Card 2: Benefits
    c2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.9), Inches(4.9))
    c2.fill.solid(); c2.fill.fore_color.rgb = LIGHT_BG
    c2.line.color.rgb = CARD_BORDER; c2.line.width = Pt(1.5)

    c2_hdr = s5.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.5), Inches(0.6))
    c2_hp = c2_hdr.text_frame.paragraphs[0]
    c2_hp.text = "• Benefits of the solution (social, economic, environmental, etc.)"
    c2_hp.font.size = Pt(15); c2_hp.font.bold = True; c2_hp.font.color.rgb = ACCENT_BLUE

    c2_body = s5.shapes.add_textbox(Inches(7.0), Inches(2.5), Inches(5.5), Inches(4.1))
    c2_btf = c2_body.text_frame; c2_btf.word_wrap = True

    benefit_items = [
        ("Social Impact: ", "Democratizes space technology; 24/7 cloud-proof situational awareness during crises."),
        ("Economic Impact: ", "Reduces GIS turnaround by 90%; eliminates expensive per-seat software licenses (ArcGIS)."),
        ("Environmental Impact: ", "Continuous, objective tracking of waterbody shrinkage, deforestation, and climate resilience."),
        ("Strategic / Atmanirbhar Bharat: ", "Native support for Indian space data (ISRO Resourcesat, Cartosat, RISAT) via Bhoonidhi."),
        ("Auditability & Trust: ", "Immutable telemetry logs with verifiable spatial masks and honest confidence scores.")
    ]

    for idx, (lead, desc) in enumerate(benefit_items):
        p = c2_btf.paragraphs[0] if idx == 0 else c2_btf.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run(); r1.text = "• " + lead; r1.font.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = NAVY
        r2 = p.add_run(); r2.text = desc; r2.font.bold = False; r2.font.size = Pt(10.5); r2.font.color.rgb = DARK_GRAY

    # Footer
    fbar = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.9), Inches(13.333), Inches(0.6))
    fbar.fill.solid(); fbar.fill.fore_color.rgb = ACCENT_BLUE; fbar.line.fill.background()

    ftxt_l = s5.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(8.0), Inches(0.5))
    ftxt_l.text_frame.text = "@SIH Idea submission- Template"
    ftxt_l.text_frame.paragraphs[0].font.color.rgb = WHITE
    ftxt_l.text_frame.paragraphs[0].font.size = Pt(12)

    ftxt_r = s5.shapes.add_textbox(Inches(12.0), Inches(6.95), Inches(1.0), Inches(0.5))
    ftxt_r.text_frame.text = "5"
    ftxt_r.text_frame.paragraphs[0].font.color.rgb = WHITE
    ftxt_r.text_frame.paragraphs[0].font.size = Pt(14)
    ftxt_r.text_frame.paragraphs[0].font.bold = True

    prs.save(pptx_path)
    print("Slide 5 appended successfully!")

if __name__ == "__main__":
    add_slide5()
